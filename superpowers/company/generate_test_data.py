#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
企业名称服务测试数据生成脚本
生成格式: json, txt, csv, xlsx, xml, sql
"""

import json
import csv
import uuid
from datetime import datetime
import xml.etree.ElementTree as ET
from xml.dom import minidom

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


def generate_uuid():
    return str(uuid.uuid4())


def generate_credit_code(prefix='91'):
    import random
    chars = '0123456789ABCDEFGHJKLMNPQRTUWXY'
    middle = ''.join(random.choice(chars) for _ in range(6))
    last = ''.join(random.choice(chars) for _ in range(10))
    return prefix + middle + last


normal_companies = [
    {
        "credit_code": "91350200MA348KJW1K",
        "name": "福建华为技术有限公司",
        "legal_rep": "任正非",
        "address": "福建省厦门市思明区软件园二期",
        "scene": "normal"
    },
    {
        "credit_code": "91310000MA1FB4HC1Q",
        "name": "上海阿里巴巴信息技术有限公司",
        "legal_rep": "张勇",
        "address": "上海市浦东新区张江高科技园区",
        "scene": "normal"
    },
    {
        "credit_code": "91440300MA5DRT3D9M",
        "name": "深圳腾讯计算机系统有限公司",
        "legal_rep": "马化腾",
        "address": "广东省深圳市南山区科技园",
        "scene": "normal"
    },
    {
        "credit_code": "91330100MA27WXK87E",
        "name": "杭州网易网络有限公司",
        "legal_rep": "丁磊",
        "address": "浙江省杭州市滨江区网易大厦",
        "scene": "normal"
    },
    {
        "credit_code": "91110000MA008T4Y8L",
        "name": "北京百度网讯科技有限公司",
        "legal_rep": "李彦宏",
        "address": "北京市海淀区中关村软件园",
        "scene": "normal"
    },
    {
        "credit_code": "91320100MA1N0T5C2P",
        "name": "南京苏宁易购电子商务有限公司",
        "legal_rep": "张近东",
        "address": "江苏省南京市玄武区苏宁大道",
        "scene": "normal"
    },
    {
        "credit_code": "91420100MA4K0X5D7Q",
        "name": "武汉小米科技有限公司",
        "legal_rep": "雷军",
        "address": "湖北省武汉市东湖新技术开发区",
        "scene": "normal"
    },
    {
        "credit_code": "91440100MA3CJC3E6H",
        "name": "广州字节跳动网络技术有限公司",
        "legal_rep": "张一鸣",
        "address": "广东省广州市天河区珠江新城",
        "scene": "normal"
    },
    {
        "credit_code": "91500000MA5U6K8E2F",
        "name": "重庆京东世纪贸易有限公司",
        "legal_rep": "刘强东",
        "address": "重庆市渝北区两江新区",
        "scene": "normal"
    },
    {
        "credit_code": "91510100MA61R8T7XQ",
        "name": "成都美团点评信息技术有限公司",
        "legal_rep": "王兴",
        "address": "四川省成都市高新区天府大道",
        "scene": "normal"
    },
    {
        "credit_code": "91350500MA2YXTK38K",
        "name": "泉州安踏体育用品有限公司",
        "legal_rep": "丁世忠",
        "address": "福建省泉州市晋江市安踏工业园",
        "scene": "normal"
    },
    {
        "credit_code": "91330200MA293KDE7Q",
        "name": "宁波雅戈尔集团股份有限公司",
        "legal_rep": "李如成",
        "address": "浙江省宁波市鄞州区雅戈尔大道",
        "scene": "normal"
    }
]

abnormal_companies = [
    {
        "credit_code": "",
        "name": "测试缺失信用代码公司",
        "legal_rep": "张三",
        "address": "北京市朝阳区",
        "scene": "missing_credit_code"
    },
    {
        "credit_code": "91350200MA348KJW1K",
        "name": "",
        "legal_rep": "李四",
        "address": "上海市浦东新区",
        "scene": "missing_name"
    },
    {
        "credit_code": "91310000MA1FB4HC1Q",
        "name": "测试缺失法定代表人公司",
        "legal_rep": "",
        "address": "广东省深圳市",
        "scene": "missing_legal_rep"
    },
    {
        "credit_code": "91350200MA348KJW1K",
        "name": "测试缺失地址公司",
        "legal_rep": "王五",
        "address": "",
        "scene": "missing_address"
    },
    {
        "credit_code": "88350200MA348KJW1K",
        "name": "测试信用代码前缀错误公司",
        "legal_rep": "赵六",
        "address": "杭州市西湖区",
        "scene": "invalid_credit_code_prefix"
    },
    {
        "credit_code": "91350200MA",
        "name": "测试信用代码长度错误公司",
        "legal_rep": "孙七",
        "address": "南京市鼓楼区",
        "scene": "invalid_credit_code_length"
    },
    {
        "credit_code": "91350200MA348KJW1K",
        "name": "ABC公司",
        "legal_rep": "周八",
        "address": "武汉市洪山区",
        "scene": "invalid_name_english"
    },
    {
        "credit_code": "91310000MA1FB4HC1Q",
        "name": "科技",
        "legal_rep": "吴九",
        "address": "广州市天河区",
        "scene": "invalid_name_too_short"
    },
    {
        "credit_code": "91440300MA5DRT3D9M",
        "name": "这是一个企业名称长度超过二十六个字符的超长名称测试公司",
        "legal_rep": "郑十",
        "address": "成都市武侯区",
        "scene": "invalid_name_too_long"
    },
    {
        "credit_code": "91330100MA27WXK87E",
        "name": "（全括号测试）",
        "legal_rep": "冯十一",
        "address": "重庆市渝中区",
        "scene": "invalid_name_all_parentheses"
    },
    {
        "credit_code": "91110000MA008T4Y8L",
        "name": "ZhangSan",
        "legal_rep": "陈十二",
        "address": "西安市雁塔区",
        "scene": "invalid_name_pinyin"
    },
    {
        "credit_code": "91320100MA1N0T5C2P",
        "name": "测试法定代表人过短公司",
        "legal_rep": "A",
        "address": "苏州市工业园区",
        "scene": "invalid_legal_rep_too_short"
    },
    {
        "credit_code": "91420100MA4K0X5D7Q",
        "name": "测试法定代表人过长公司",
        "legal_rep": "一二三四五六七",
        "address": "青岛市崂山区",
        "scene": "invalid_legal_rep_too_long"
    },
    {
        "credit_code": "91440100MA3CJC3E6H",
        "name": "测试法定代表人非汉字公司",
        "legal_rep": "LiMing",
        "address": "天津市南开区",
        "scene": "invalid_legal_rep_not_chinese"
    },
    {
        "credit_code": "91350200MA348KJW1K",
        "name": "同名不同企测试公司",
        "legal_rep": "黄十五",
        "address": "长沙市岳麓区",
        "scene": "same_name_different_company_1"
    },
    {
        "credit_code": "91310000MA1FB4HC2Q",
        "name": "同名不同企测试公司",
        "legal_rep": "杨十六",
        "address": "郑州市金水区",
        "scene": "same_name_different_company_2"
    },
    {
        "credit_code": "91440300MA5DRT3D9M",
        "name": "同企不同名测试公司（旧名）",
        "legal_rep": "朱十七",
        "address": "合肥市蜀山区",
        "scene": "same_company_different_name_1"
    },
    {
        "credit_code": "91440300MA5DRT3D9M",
        "name": "同企不同名测试公司（新名）",
        "legal_rep": "朱十七",
        "address": "合肥市蜀山区",
        "scene": "same_company_different_name_2"
    }
]

all_companies = normal_companies + abnormal_companies


def save_json(companies, filepath):
    data = {
        "generated_at": datetime.now().isoformat(),
        "total_count": len(companies),
        "normal_count": len([c for c in companies if c['scene'] == 'normal']),
        "abnormal_count": len([c for c in companies if c['scene'] != 'normal']),
        "companies": [
            {
                "ent_code": generate_uuid(),
                "credit_code": c['credit_code'],
                "name": c['name'],
                "legal_rep": c['legal_rep'],
                "address": c['address'],
                "scene": c['scene']
            }
            for c in companies
        ]
    }
    with open(filepath, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    print(f"JSON文件已保存: {filepath}")


def save_txt(companies, filepath):
    lines = []
    lines.append("=" * 80)
    lines.append("企业名称服务测试数据")
    lines.append(f"生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    lines.append(f"总记录数: {len(companies)}")
    lines.append(f"正常数据: {len([c for c in companies if c['scene'] == 'normal'])}")
    lines.append(f"异常数据: {len([c for c in companies if c['scene'] != 'normal'])}")
    lines.append("=" * 80)
    lines.append("")
    
    lines.append("【正常企业数据】")
    lines.append("-" * 80)
    for i, c in enumerate([c for c in companies if c['scene'] == 'normal'], 1):
        lines.append(f"{i}. 统一社会信用代码: {c['credit_code']}")
        lines.append(f"   企业名称: {c['name']}")
        lines.append(f"   法定代表人: {c['legal_rep']}")
        lines.append(f"   地址: {c['address']}")
        lines.append("")
    
    lines.append("【异常企业数据】")
    lines.append("-" * 80)
    for i, c in enumerate([c for c in companies if c['scene'] != 'normal'], 1):
        lines.append(f"{i}. 场景类型: {c['scene']}")
        lines.append(f"   统一社会信用代码: {c['credit_code']}")
        lines.append(f"   企业名称: {c['name']}")
        lines.append(f"   法定代表人: {c['legal_rep']}")
        lines.append(f"   地址: {c['address']}")
        lines.append("")
    
    with open(filepath, 'w', encoding='utf-8') as f:
        f.write('\n'.join(lines))
    print(f"TXT文件已保存: {filepath}")


def save_csv(companies, filepath):
    with open(filepath, 'w', encoding='utf-8-sig', newline='') as f:
        writer = csv.writer(f)
        writer.writerow(['序号', '统一社会信用代码', '企业名称', '法定代表人', '地址', '场景类型'])
        for i, c in enumerate(companies, 1):
            writer.writerow([i, c['credit_code'], c['name'], c['legal_rep'], c['address'], c['scene']])
    print(f"CSV文件已保存: {filepath}")


def save_xlsx(companies, filepath):
    if not HAS_OPENPYXL:
        print("警告: openpyxl未安装，跳过XLSX文件生成")
        print("可执行: pip install openpyxl")
        return
    
    wb = Workbook()
    
    ws1 = wb.active
    ws1.title = "全部数据"
    ws1.append(['序号', '统一社会信用代码', '企业名称', '法定代表人', '地址', '场景类型'])
    for i, c in enumerate(companies, 1):
        ws1.append([i, c['credit_code'], c['name'], c['legal_rep'], c['address'], c['scene']])
    
    ws2 = wb.create_sheet("正常数据")
    ws2.append(['序号', '统一社会信用代码', '企业名称', '法定代表人', '地址'])
    for i, c in enumerate([c for c in companies if c['scene'] == 'normal'], 1):
        ws2.append([i, c['credit_code'], c['name'], c['legal_rep'], c['address']])
    
    ws3 = wb.create_sheet("异常数据")
    ws3.append(['序号', '异常类型', '统一社会信用代码', '企业名称', '法定代表人', '地址'])
    for i, c in enumerate([c for c in companies if c['scene'] != 'normal'], 1):
        ws3.append([i, c['scene'], c['credit_code'], c['name'], c['legal_rep'], c['address']])
    
    for ws in [ws1, ws2, ws3]:
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal='center')
    
    wb.save(filepath)
    print(f"XLSX文件已保存: {filepath}")


def save_xml(companies, filepath):
    root = ET.Element('company_test_data')
    root.set('generated_at', datetime.now().isoformat())
    root.set('total_count', str(len(companies)))
    
    normal = ET.SubElement(root, 'normal_companies')
    normal.set('count', str(len([c for c in companies if c['scene'] == 'normal'])))
    for c in [c for c in companies if c['scene'] == 'normal']:
        company = ET.SubElement(normal, 'company')
        ET.SubElement(company, 'ent_code').text = generate_uuid()
        ET.SubElement(company, 'credit_code').text = c['credit_code']
        ET.SubElement(company, 'name').text = c['name']
        ET.SubElement(company, 'legal_rep').text = c['legal_rep']
        ET.SubElement(company, 'address').text = c['address']
    
    abnormal = ET.SubElement(root, 'abnormal_companies')
    abnormal.set('count', str(len([c for c in companies if c['scene'] != 'normal'])))
    for c in [c for c in companies if c['scene'] != 'normal']:
        company = ET.SubElement(abnormal, 'company')
        company.set('scene', c['scene'])
        ET.SubElement(company, 'credit_code').text = c['credit_code']
        ET.SubElement(company, 'name').text = c['name']
        ET.SubElement(company, 'legal_rep').text = c['legal_rep']
        ET.SubElement(company, 'address').text = c['address']
    
    rough_string = ET.tostring(root, 'utf-8')
    reparsed = minidom.parseString(rough_string)
    pretty_xml = reparsed.toprettyxml(indent="  ", encoding='utf-8')
    
    with open(filepath, 'wb') as f:
        f.write(pretty_xml)
    print(f"XML文件已保存: {filepath}")


def save_sql(companies, filepath):
    lines = []
    lines.append("-- =============================================")
    lines.append("-- 企业名称服务测试数据SQL脚本")
    lines.append(f"-- 生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    lines.append("-- =============================================")
    lines.append("")
    
    lines.append("-- 插入正常企业数据")
    lines.append("INSERT INTO company_fact (ent_code, credit_code, name, legal_rep, address, ent_record_scene, create_time) VALUES")
    values = []
    for c in [c for c in companies if c['scene'] == 'normal']:
        ent_code = generate_uuid()
        values.append(f"('{ent_code}', '{c['credit_code']}', '{c['name']}', '{c['legal_rep']}', '{c['address']}', 1, CURRENT_TIMESTAMP)")
    lines.append(',\n'.join(values) + ';')
    lines.append("")
    
    lines.append("-- 插入异常企业数据（用于测试校验逻辑）")
    lines.append("-- 注意：这些数据可能违反约束，仅用于测试校验逻辑")
    lines.append("/*")
    for c in [c for c in companies if c['scene'] != 'normal']:
        ent_code = generate_uuid()
        lines.append(f"-- 场景: {c['scene']}")
        lines.append(f"-- INSERT INTO company_fact (ent_code, credit_code, name, legal_rep, address, ent_record_scene)")
        lines.append(f"-- VALUES ('{ent_code}', '{c['credit_code']}', '{c['name']}', '{c['legal_rep']}', '{c['address']}', 1);")
        lines.append("")
    lines.append("*/")
    
    with open(filepath, 'w', encoding='utf-8') as f:
        f.write('\n'.join(lines))
    print(f"SQL文件已保存: {filepath}")


def main():
    base_dir = r"D:\code-1\xiaoyi-trae\eiker-be\docs\superpowers\company"
    
    save_json(all_companies, f"{base_dir}\\company_test_data.json")
    save_txt(all_companies, f"{base_dir}\\company_test_data.txt")
    save_csv(all_companies, f"{base_dir}\\company_test_data.csv")
    save_xlsx(all_companies, f"{base_dir}\\company_test_data.xlsx")
    save_xml(all_companies, f"{base_dir}\\company_test_data.xml")
    save_sql(all_companies, f"{base_dir}\\company_test_data_insert.sql")
    
    print("")
    print("=" * 50)
    print(f"测试数据生成完成！")
    print(f"总记录数: {len(all_companies)}")
    print(f"正常数据: {len(normal_companies)}")
    print(f"异常数据: {len(abnormal_companies)}")
    print("=" * 50)


if __name__ == '__main__':
    main()
